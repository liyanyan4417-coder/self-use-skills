#!/usr/bin/env python3
import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from scipy.stats import chi2 as chi2_dist
from semopy import Model, calc_stats


def read_data(path):
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    raise ValueError(f"Unsupported data file: {path}")


def cronbach_alpha(frame):
    clean = frame.dropna()
    k = clean.shape[1]
    item_variances = clean.var(axis=0, ddof=1)
    total_variance = clean.sum(axis=1).var(ddof=1)
    return k / (k - 1) * (1 - item_variances.sum() / total_variance)


def std_col(estimates):
    for column in ["Est. Std", "Std. Est", "Std.Estimate"]:
        if column in estimates.columns:
            return column
    raise ValueError(f"No standardized estimate column found: {list(estimates.columns)}")


def p_marker(value, decimals):
    if value in ("", "-", None) or pd.isna(value):
        return "-"
    value = float(value)
    if value < 0.001:
        return "***"
    if value < 0.01:
        return "**"
    if value < 0.05:
        return "*"
    return round(value, decimals)


def round_value(value, decimals):
    if value in ("", "-", None):
        return value
    if isinstance(value, str):
        return value
    if pd.isna(value):
        return ""
    return round(float(value), decimals)


def model_syntax(constructs, structural_paths=None, measurement_only=False):
    lines = []
    for construct, spec in constructs.items():
        lines.append(f"{construct} =~ {' + '.join(spec['items'])}")
    lines.append("")

    names = list(constructs)
    if measurement_only:
        for i, left in enumerate(names):
            for right in names[i + 1 :]:
                lines.append(f"{left} ~~ {right}")
    else:
        structural_paths = structural_paths or []
        endogenous = set()
        for path in structural_paths:
            dep = path["dependent"]
            preds = path["predictors"]
            endogenous.add(dep)
            lines.append(f"{dep} ~ {' + '.join(preds)}")
        lines.append("")
        exogenous = [name for name in names if name not in endogenous]
        for i, left in enumerate(exogenous):
            for right in exogenous[i + 1 :]:
                lines.append(f"{left} ~~ {right}")
    return "\n".join(lines)


def fit_model(desc, data):
    model = Model(desc)
    model.fit(data)
    return model, model.inspect(std_est=True), calc_stats(model).T


def measurement_loadings(estimates, constructs):
    sc = std_col(estimates)
    observed = {item for spec in constructs.values() for item in spec["items"]}
    rows = estimates.loc[estimates["op"].eq("~")].copy()
    rows = rows.loc[rows["rval"].isin(constructs) & rows["lval"].isin(observed)]
    rows = rows[["rval", "lval", "Estimate", sc, "Std. Err", "z-value", "p-value"]]
    rows.columns = [
        "construct",
        "item",
        "unstandardized_loading",
        "standardized_loading",
        "std_error",
        "z_value",
        "p_value",
    ]
    rows["construct_label"] = rows["construct"].map(lambda x: constructs[x].get("label", x))
    rows["SMC"] = rows["standardized_loading"].astype(float) ** 2
    return rows.sort_values(["construct", "item"]).reset_index(drop=True)


def reliability_summary(data, constructs, loadings):
    rows = []
    for construct, spec in constructs.items():
        items = spec["items"]
        lambdas = loadings.loc[loadings["construct"].eq(construct), "standardized_loading"].astype(float).to_numpy()
        errors = 1 - np.square(lambdas)
        cr = np.square(lambdas.sum()) / (np.square(lambdas.sum()) + errors.sum())
        ave = np.square(lambdas).sum() / (np.square(lambdas).sum() + errors.sum())
        rows.append(
            {
                "construct": construct,
                "label": spec.get("label", construct),
                "items": ", ".join(items),
                "alpha": cronbach_alpha(data[items]),
                "CR": cr,
                "AVE": ave,
                "sqrt_AVE": np.sqrt(ave),
            }
        )
    return pd.DataFrame(rows)


def latent_correlation_matrix(estimates, constructs, reliability):
    sc = std_col(estimates)
    names = list(constructs)
    corr = pd.DataFrame(np.eye(len(names)), index=names, columns=names)
    rows = estimates.loc[
        estimates["op"].eq("~~") & estimates["lval"].isin(names) & estimates["rval"].isin(names)
    ]
    for _, row in rows.iterrows():
        if row["lval"] != row["rval"]:
            corr.loc[row["lval"], row["rval"]] = row[sc]
            corr.loc[row["rval"], row["lval"]] = row[sc]
    sqrt_ave = reliability.set_index("construct")["sqrt_AVE"]
    for name in names:
        corr.loc[name, name] = sqrt_ave.loc[name]
    labels = [constructs[name].get("label", name) for name in names]
    corr.index = labels
    corr.columns = labels
    return corr


def htmt_matrix(data, constructs):
    names = list(constructs)
    corr = data.corr().abs()
    htmt = pd.DataFrame(np.eye(len(names)), index=names, columns=names)
    for i, a in enumerate(names):
        for b in names[i + 1 :]:
            items_a = constructs[a]["items"]
            items_b = constructs[b]["items"]
            heterotrait = corr.loc[items_a, items_b].to_numpy().mean()
            monotrait_a = corr.loc[items_a, items_a].where(~np.eye(len(items_a), dtype=bool)).stack().mean()
            monotrait_b = corr.loc[items_b, items_b].where(~np.eye(len(items_b), dtype=bool)).stack().mean()
            value = heterotrait / np.sqrt(monotrait_a * monotrait_b)
            htmt.loc[a, b] = value
            htmt.loc[b, a] = value
    labels = [constructs[name].get("label", name) for name in names]
    htmt.index = labels
    htmt.columns = labels
    return htmt


def srmr(model, data):
    observed = model.vars["observed"]
    sample_corr = data[observed].corr().to_numpy()
    implied_cov = model.calc_sigma()[0]
    implied_std = np.sqrt(np.diag(implied_cov))
    implied_corr = implied_cov / np.outer(implied_std, implied_std)
    residual = sample_corr - implied_corr
    lower = residual[np.tril_indices_from(residual, k=-1)]
    return float(np.sqrt(np.mean(lower**2)))


def fit_rows(model, data, fit_stats):
    values = fit_stats["Value"]
    chi_square = float(values["chi2"])
    df = float(values["DoF"])
    baseline_chi_square = float(values["chi2 Baseline"])
    n = len(data)
    p = data.shape[1]
    ifi = (baseline_chi_square - chi_square) / (baseline_chi_square - df)
    hoelter_n = ((chi2_dist.ppf(0.95, df) / chi_square) * (n - 1)) + 1
    gamma_hat = p / (p + 2 * max(chi_square - df, 0) / (n - 1))
    mcdonald_nci = np.exp(-max(chi_square - df, 0) / (2 * (n - 1)))
    return [
        ("卡方值 (Chi-square)", "越小越好", chi_square),
        ("自由度 (df)", "越大越好", df),
        ("x²/df", "1< x²/df< 3.0", chi_square / df),
        ("拟合优度指数 (GFI)", "> 0.90", values["GFI"]),
        ("调整后拟合优度指数 (AGFI)", "> 0.90", values["AGFI"]),
        ("近似误差均方根 (RMSEA)", "< 0.08", values["RMSEA"]),
        ("标准化残差均方根 (SRMR)", "< 0.08", srmr(model, data)),
        ("比较拟合指数 (CFI)", "> 0.90", values["CFI"]),
        ("IFI", "> 0.90", ifi),
        ("非规范拟合指数 (TLI)", "> 0.90", values["TLI"]),
        ("Hoelter's N", ">200", hoelter_n),
        ("Gamma hat", "> 0.90", gamma_hat),
        ("McDonald's NCI", ">0.9", mcdonald_nci),
    ]


def set_border(ws, min_row, max_row, min_col, max_col, top=False, bottom=False):
    top_side = Side(style="medium" if top else None, color="000000")
    bottom_side = Side(style="medium" if bottom else None, color="000000")
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(top=top_side, bottom=bottom_side)


def write_table1(ws, constructs, loadings, reliability, decimals):
    ws["A1"] = "表1  参数估计和构面信效度"
    ws["A1"].font = Font(name="Songti SC", bold=True, size=14)
    ws.merge_cells("A1:J1")
    for region in ["A2:A3", "B2:B3", "C2:F2", "G2:H2", "I2:I3", "J2:J3"]:
        ws.merge_cells(region)
    headers = {
        "A2": "构面",
        "B2": "题目",
        "C2": "参数显著性估计",
        "G2": "题目信度",
        "I2": "组成\n信度",
        "J2": "收敛\n效度",
        "C3": "Unstd.",
        "D3": "S.E.",
        "E3": "z-value",
        "F3": "P",
        "G3": "Std.",
        "H3": "SMC",
    }
    for cell, value in headers.items():
        ws[cell] = value
    row = 4
    rel = reliability.set_index("construct")
    for construct, spec in constructs.items():
        item_rows = loadings.set_index("item").loc[spec["items"]].reset_index()
        start = row
        for idx, item in enumerate(spec["items"]):
            item_row = item_rows.loc[item_rows["item"].eq(item)].iloc[0]
            ws.cell(row=row, column=2, value=item)
            ws.cell(row=row, column=3, value=round_value(item_row["unstandardized_loading"], decimals))
            ws.cell(row=row, column=4, value=round_value(item_row["std_error"], decimals))
            ws.cell(row=row, column=5, value=round_value(item_row["z_value"], decimals))
            ws.cell(row=row, column=6, value=p_marker(item_row["p_value"], decimals))
            ws.cell(row=row, column=7, value=round_value(item_row["standardized_loading"], decimals))
            ws.cell(row=row, column=8, value=round_value(item_row["SMC"], decimals))
            if idx == 0:
                ws.cell(row=row, column=1, value=spec.get("label", construct))
                ws.cell(row=row, column=9, value=round_value(rel.loc[construct, "CR"], decimals))
                ws.cell(row=row, column=10, value=round_value(rel.loc[construct, "AVE"], decimals))
            row += 1
        ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)
    set_border(ws, 1, 1, 1, 10, bottom=True)
    set_border(ws, 3, 3, 1, 10, bottom=True)
    set_border(ws, row - 1, row - 1, 1, 10, bottom=True)
    for index, width in enumerate([14, 10, 12, 12, 12, 10, 10, 10, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    return row + 2


def write_table2(ws, start, constructs, reliability, fornell, htmt, decimals):
    ws.cell(row=start, column=1, value="表2区分效度")
    ws.cell(row=start, column=1).font = Font(name="Songti SC", bold=True, size=14)
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=7)
    set_border(ws, start, start, 1, 7, bottom=True)
    ws.merge_cells(start_row=start + 1, start_column=2, end_row=start + 1, end_column=2)
    ws.merge_cells(start_row=start + 1, start_column=3, end_row=start + 1, end_column=6)
    ws.cell(row=start + 1, column=2, value="收敛效度")
    ws.cell(row=start + 1, column=3, value="区分效度")
    ws.cell(row=start + 2, column=2, value="AVE")
    labels = [spec.get("label", name) for name, spec in constructs.items()]
    for idx, label in enumerate(labels, start=3):
        ws.cell(row=start + 2, column=idx, value=label)
    rel = reliability.set_index("construct")
    for row_idx, (construct, spec) in enumerate(constructs.items(), start=start + 3):
        label = spec.get("label", construct)
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=round_value(rel.loc[construct, "AVE"], decimals))
        for col_idx, col_label in enumerate(labels, start=3):
            row_pos = row_idx - (start + 3)
            col_pos = col_idx - 3
            if row_pos == col_pos:
                cell = ws.cell(row=row_idx, column=col_idx, value=round_value(fornell.loc[label, col_label], decimals))
                cell.font = Font(name="Songti SC", bold=True)
            elif row_pos > col_pos:
                ws.cell(row=row_idx, column=col_idx, value=round_value(fornell.loc[label, col_label], decimals))
            else:
                ws.cell(row=row_idx, column=col_idx, value=round_value(htmt.loc[label, col_label], decimals))
    note_row = start + 3 + len(constructs)
    ws.cell(
        row=note_row,
        column=1,
        value="注：对角线粗体字为AVE之开根号值，下三角为构面皮尔森相关，上三角为HTMT异质特质－同质特质相关比",
    )
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    set_border(ws, start + 1, start + 1, 2, 6, bottom=True)
    set_border(ws, start + 2, start + 2, 1, 6, bottom=True)
    return note_row + 2


def write_table3(ws, start, rows, decimals):
    ws.cell(row=start, column=1, value="表3  模型拟合度")
    ws.cell(row=start, column=1).font = Font(name="Songti SC", bold=True, size=14)
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=3)
    set_border(ws, start, start, 1, 3, bottom=True)
    for col, header in enumerate(["拟合指标（Index）", "评价标准\n(Criteria)", "分析结果\n(Result)"], start=1):
        ws.cell(row=start + 1, column=col, value=header)
    for offset, (index, criteria, result) in enumerate(rows, start=2):
        ws.cell(row=start + offset, column=1, value=index)
        ws.cell(row=start + offset, column=2, value=criteria)
        ws.cell(row=start + offset, column=3, value=round_value(result, decimals))
    set_border(ws, start + 1, start + 1, 1, 3, bottom=True)
    set_border(ws, start + 1 + len(rows), start + 1 + len(rows), 1, 3, bottom=True)
    for column, width in {"A": 28, "B": 22, "C": 16}.items():
        ws.column_dimensions[column].width = max(ws.column_dimensions[column].width or 0, width)


def style_sheet(ws, decimals):
    for row in ws.iter_rows():
        for cell in row:
            cell.font = cell.font.copy(name="Songti SC", size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0." + ("0" * decimals)
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 23


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True)
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text(encoding="utf-8"))

    decimals = int(config.get("decimals", 3))
    constructs = config["constructs"]
    output_file = Path(config.get("output_file", "outputs/sem_measurement_model_tables.xlsx"))
    output_file.parent.mkdir(parents=True, exist_ok=True)

    data = read_data(config["data_file"])
    items = [item for spec in constructs.values() for item in spec["items"]]
    missing = [item for item in items if item not in data.columns]
    if missing:
        raise ValueError(f"Missing item columns in data: {missing}")
    data_for_sem = data[items].copy()

    measurement_desc = model_syntax(constructs, measurement_only=True)
    measurement_model, measurement_estimates, measurement_fit = fit_model(measurement_desc, data_for_sem)
    loadings = measurement_loadings(measurement_estimates, constructs)
    reliability = reliability_summary(data_for_sem, constructs, loadings)
    fornell = latent_correlation_matrix(measurement_estimates, constructs, reliability)
    htmt = htmt_matrix(data_for_sem, constructs)
    fits = fit_rows(measurement_model, data_for_sem, measurement_fit)

    wb = Workbook()
    ws = wb.active
    ws.title = "SEM三表"
    next_row = write_table1(ws, constructs, loadings, reliability, decimals)
    next_row = write_table2(ws, next_row, constructs, reliability, fornell, htmt, decimals)
    write_table3(ws, next_row, fits, decimals)
    style_sheet(ws, decimals)

    if config.get("structural_paths"):
        structural_desc = model_syntax(constructs, config["structural_paths"], measurement_only=False)
        _, structural_estimates, _ = fit_model(structural_desc, data_for_sem)
        appendix = wb.create_sheet("Structural_Estimates")
        for col_idx, name in enumerate(structural_estimates.columns, start=1):
            appendix.cell(row=1, column=col_idx, value=name)
        for row_idx, row in enumerate(structural_estimates.itertuples(index=False), start=2):
            for col_idx, value in enumerate(row, start=1):
                appendix.cell(row=row_idx, column=col_idx, value=round_value(value, decimals))

    wb.save(output_file)
    print(output_file)


if __name__ == "__main__":
    main()
